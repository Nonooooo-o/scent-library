"""Import exact supplied workbook values; cross-check every rating and image against Word."""
import argparse, collections, hashlib, json, re
from pathlib import Path
import openpyxl
from docx import Document
from PIL import Image
parser=argparse.ArgumentParser();parser.add_argument('xlsx');parser.add_argument('docx');args=parser.parse_args()
root=Path(__file__).resolve().parent.parent
old=json.loads((root/'data/perfumes.json').read_text())
wb=openpyxl.load_workbook(args.xlsx,data_only=True);ws=wb.worksheets[0]
assert ws.max_row-4==len(old)==1612
images={im.anchor._from.row+1:im._data() for im in ws._images}
assert len(images)==1612
word=[]
for p in Document(args.docx).paragraphs:
 m=re.fullmatch(r'([1-5])\s*/\s*5',p.text.strip())
 if m:
  b=p._p.xpath('.//a:blip');assert len(b)==1
  rid=b[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
  word.append((int(m.group(1)),p.part.rels[rid].target_part.blob))
assert len(word)==1612
checks=[]
for i,p in enumerate(old):
 row=i+5;brand,zh,en,year,rating,desc=[ws.cell(row,j).value for j in range(2,8)]
 year=None if year=='待核实' else year
 assert [p['brand'],p['nameChinese'],p['nameEnglish'],p['releaseYear'],p['personalRating']]==[brand,zh,en,year,rating]
 assert word[i][0]==rating
 blob=images[row];assert blob==word[i][1]
 original=(root/'public'/p['originalImage'].lstrip('/')).read_bytes();assert blob==original
 p['description']=desc or ''
 m=re.search(r'^香调轮廓以(.+?)为主',desc or '')
 p['scents']=m.group(1).split('、') if m else []
 p['image']=p['imageThumbnail']
 im=Image.open(root/'public'/p['image'].lstrip('/'))
 p['imageWidth'],p['imageHeight']=im.size
 p['sourceOrder']=i+1
 p['yearLabel']=str(year) if year is not None else '待核实'
 p['source']={'file':'香水(1).xlsx','sheet':ws.title,'row':row}
 # Keep prior research in its existing fields for historical traceability;
 # the new UI uses only exact workbook description and parsed scents.
 checks.append({'id':p['id'],'row':row,'sha256':hashlib.sha256(blob).hexdigest(),'image':p['image']})
(root/'data/perfumes.json').write_text(json.dumps(old,ensure_ascii=False,indent=2)+'\n')
fields=['id','brand','brandSlug','nameChinese','nameEnglish','releaseYear','personalRating','description','scents','image','imageWidth','imageHeight','sourceOrder','yearLabel']
(root/'data/collection.json').write_text(json.dumps([{k:p[k] for k in fields} for p in old],ensure_ascii=False,separators=(',',':'))+'\n')
audit={'records':len(old),'brands':len(set(p['brand'] for p in old)),'personalRatingsMatchedToWord':1612,'imagesMatchedToWordAndExcel':1612,'stableIdsPreserved':1612,'missingYears':sum(p['releaseYear'] is None for p in old),'ratings':dict(collections.Counter(p['personalRating'] for p in old)),'descriptionSource':'Exact uploaded workbook values','scentParsing':'Only literal terms in 香调轮廓以…为主','checks':checks}
(root/'data/source-audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:v for k,v in audit.items() if k!='checks'},ensure_ascii=False))
